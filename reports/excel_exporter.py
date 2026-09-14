"""Excel report generation with openpyxl.

Produces a professional multi-sheet workbook: frozen header rows, auto-filters,
sensible column widths, currency formatting, clickable hyperlinks and a
generated timestamp. (openpyxl is used directly rather than pandas for precise
control over formatting and hyperlinks.)
"""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import config
from services import tender_service as ts
from utils.dates import now_tz, days_until

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_LINK_FONT = Font(color="0563C1", underline="single")
_TITLE_FONT = Font(bold=True, size=13, color="1F4E78")
_ALIGN = Alignment(vertical="top", wrap_text=False)
_THIN = Side(style="thin", color="D9D9D9")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

# Column spec for tender sheets: (header, key, kind)
_TENDER_COLUMNS = [
    ("Portal", "portal_name", "text"),
    ("Tender ID", "tender_id", "text"),
    ("Reference", "tender_reference", "text"),
    ("Title", "tender_title", "text"),
    ("Organisation", "organisation", "text"),
    ("Category", "tender_category", "text"),
    ("Location", "location", "text"),
    ("Published", "published_date", "text"),
    ("Closing", "bid_submission_end", "text"),
    ("Opening", "opening_date", "text"),
    ("Estimated Value", "estimated_value", "currency"),
    ("EMD", "emd", "text"),
    ("Tender Fee", "tender_fee", "text"),
    ("Relevance", "relevance_band", "text"),
    ("Score", "relevance_score", "text"),
    ("Status", "status", "text"),
    ("Detail URL", "tender_detail_url", "url"),
    ("Documents", "document_urls", "doclist"),
]


def _cell_value(row, key, kind):
    value = row.get(key)
    if kind == "doclist":
        return " | ".join(value or []) if isinstance(value, list) else (value or "")
    if kind == "currency":
        return value if isinstance(value, (int, float)) else None
    return "" if value is None else value


def _write_sheet(wb, title, columns, rows):
    ws = wb.create_sheet(title[:31])
    headers = [c[0] for c in columns]
    ws.append(headers)
    for c_idx, cell in enumerate(ws[1], start=1):
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        cell.border = _BORDER

    for row in rows:
        values = [_cell_value(row, key, kind) for _h, key, kind in columns]
        ws.append(values)
        r = ws.max_row
        for idx, (_h, key, kind) in enumerate(columns, start=1):
            cell = ws.cell(row=r, column=idx)
            cell.alignment = _ALIGN
            cell.border = _BORDER
            if kind == "currency" and isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0"
            elif kind == "url" and cell.value:
                cell.hyperlink = cell.value
                cell.font = _LINK_FONT
            elif kind == "localpath" and cell.value:
                cell.hyperlink = "file:///" + str(cell.value).replace("\\", "/")
                cell.font = _LINK_FONT

    _finalize(ws, columns)
    return ws


def _finalize(ws, columns):
    last_col = get_column_letter(len(columns))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:%s%d" % (last_col, max(ws.max_row, 1))
    for idx, (header, key, kind) in enumerate(columns, start=1):
        letter = get_column_letter(idx)
        max_len = len(str(header))
        for cell in ws[letter]:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max(10, max_len + 2), 60)


def _is_active(tender, tz):
    d = days_until(tender.get("bid_submission_end"), tz)
    return d is None or d >= 0


def _closing_within(tenders, days, tz):
    result = []
    for t in tenders:
        d = days_until(t.get("bid_submission_end"), tz)
        if d is not None and 0 <= d <= days:
            result.append(t)
    return result


def generate_report(path=None):
    tz = ts.get_timezone_name()
    now = now_tz(tz)
    if path is None:
        config.ensure_directories()
        path = config.REPORTS_DIR / ("Tender_Report_%s.xlsx" % now.strftime("%Y-%m-%d"))
    path = str(path)

    all_rows = ts.all_tenders()
    active = [t for t in all_rows if _is_active(t, tz)]
    new_rows = [t for t in all_rows if t.get("status") == "new"]
    updated_rows = [t for t in all_rows if t.get("status") == "updated"]

    portals = ts.list_portals()
    for p in portals:
        p["enabled_label"] = "Yes" if p.get("enabled") else "No"
        p["tender_count"] = ts.count_tenders_for_portal(p["id"])

    wb = Workbook()
    wb.remove(wb.active)

    _write_sheet(wb, "NEW TENDERS", _TENDER_COLUMNS, new_rows)
    _write_sheet(wb, "UPDATED TENDERS", _TENDER_COLUMNS, updated_rows)
    _write_sheet(wb, "ALL ACTIVE TENDERS", _TENDER_COLUMNS, active)
    _write_sheet(wb, "CLOSING WITHIN 7 DAYS", _TENDER_COLUMNS, _closing_within(active, 7, tz))
    _write_sheet(wb, "CLOSING WITHIN 30 DAYS", _TENDER_COLUMNS, _closing_within(active, 30, tz))

    _write_sheet(wb, "DOCUMENT DOWNLOADS", [
        ("Portal", "portal_name", "text"),
        ("Tender", "tender_title", "text"),
        ("Document", "document_name", "text"),
        ("Status", "status", "text"),
        ("Size (bytes)", "file_size", "currency"),
        ("Downloaded At", "downloaded_at", "text"),
        ("Document URL", "document_url", "url"),
        ("Local Path", "local_path", "localpath"),
    ], ts.all_documents())

    _write_sheet(wb, "PORTAL STATUS", [
        ("Name", "name", "text"),
        ("URL", "url", "url"),
        ("Enabled", "enabled_label", "text"),
        ("Schedule", "schedule_type", "text"),
        ("Run Time", "run_time", "text"),
        ("Health", "health", "text"),
        ("Last Status", "last_status", "text"),
        ("Last Scan", "last_scan_at", "text"),
        ("Tenders", "tender_count", "currency"),
    ], portals)

    scan_sheet = _write_sheet(wb, "SCAN SUMMARY", [
        ("Started", "started_at", "text"),
        ("Portal", "portal_name", "text"),
        ("Duration (s)", "duration_seconds", "currency"),
        ("Pages", "pages_processed", "currency"),
        ("Found", "tenders_found", "currency"),
        ("New", "new_count", "currency"),
        ("Updated", "updated_count", "currency"),
        ("Existing", "existing_count", "currency"),
        ("Documents", "documents_downloaded", "currency"),
        ("Errors", "error_count", "currency"),
        ("Status", "status", "text"),
    ], ts.list_scans(500))

    _write_sheet(wb, "ERRORS", [
        ("Time", "occurred_at", "text"),
        ("Portal", "portal_name", "text"),
        ("Type", "error_type", "text"),
        ("Message", "message", "text"),
    ], ts.recent_errors(1000))

    # Generated timestamp (workbook properties + a note below the scan table).
    wb.properties.created = now.replace(tzinfo=None)
    note_row = scan_sheet.max_row + 2
    scan_sheet.cell(row=note_row, column=1, value="Report generated:").font = _TITLE_FONT
    scan_sheet.cell(row=note_row, column=2, value=now.strftime("%Y-%m-%d %H:%M:%S %Z"))

    wb.save(path)
    return path
