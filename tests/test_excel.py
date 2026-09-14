"""Excel report generation."""
from openpyxl import load_workbook

from services import tender_service as ts
from database.models import new_tender_dict
from reports import excel_exporter

_EXPECTED_SHEETS = [
    "NEW TENDERS", "UPDATED TENDERS", "ALL ACTIVE TENDERS",
    "CLOSING WITHIN 7 DAYS", "CLOSING WITHIN 30 DAYS", "DOCUMENT DOWNLOADS",
    "PORTAL STATUS", "SCAN SUMMARY", "ERRORS",
]


def _add_tender(portal, tid, title, status, closing="2026-12-31", value=500000):
    t = new_tender_dict(portal["id"])
    t.update({"tender_id": tid, "tender_title": title, "portal_name": portal["name"],
              "organisation": "PHED", "bid_submission_end": closing,
              "estimated_value": value, "estimated_value_text": "Rs. 5 Lakh",
              "relevance_band": "HIGH", "relevance_score": 80})
    t["unique_key"] = ts.compute_unique_key(t)
    return ts.insert_tender(t, status)


def test_generate_report(db, sample_portal, tmp_path):
    _add_tender(sample_portal, "T-1", "Water new", "new")
    _add_tender(sample_portal, "T-2", "Water updated", "updated")
    _add_tender(sample_portal, "T-3", "Water existing", "existing")
    ts.create_scan(sample_portal["id"], sample_portal["name"], "manual")

    path = excel_exporter.generate_report(str(tmp_path / "report.xlsx"))
    wb = load_workbook(path)
    for sheet in _EXPECTED_SHEETS:
        assert sheet in wb.sheetnames

    new_ws = wb["NEW TENDERS"]
    assert new_ws["A1"].value == "Portal"          # header row
    assert new_ws.freeze_panes == "A2"             # header frozen
    assert new_ws.auto_filter.ref is not None      # filters enabled
    titles = [new_ws.cell(row=r, column=4).value for r in range(2, new_ws.max_row + 1)]
    assert "Water new" in titles


def test_report_default_path(db, sample_portal):
    _add_tender(sample_portal, "T-9", "Water", "new")
    path = excel_exporter.generate_report()
    assert path.endswith(".xlsx")
    import os
    assert os.path.exists(path)
